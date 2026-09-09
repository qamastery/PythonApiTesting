import json
import jsonpath
import requests
import openpyxl

class Common:

    def __init__(self, FileNamePath, SheetName):
        global wb
        global sh
        wb = openpyxl.load_workbook(FileNamePath)
        sh = wb[SheetName]

    def fetch_row_count(self):
        rows = sh.max_row
        return rows

    def fetch_column_count(self):
        columns = sh.max_column
        return columns

    def fetch_key_names(self):
        c = sh.max_column
        list=[]
        for i in range (1, c+1):
            cell = sh.cell(row=1,column=i)
            list.insert(i-1,cell.value)

        return list

    def update_request_with_data(self,rowNumber,jsonRequest,keyList):
       c = sh.max_column       # 4
       for i in range(1, c+1):
           cell = sh.cell(row=rowNumber, column=i)
           jsonRequest[keyList[i-1]]=cell.value

       return jsonRequest