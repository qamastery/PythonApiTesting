import requests
import json
import jsonpath
import openpyxl

def test_add_multiple_students():
    # API
    API_URL = "https://thetestingworldapi.com/api/studentsDetails"
    file = open('/home/qamaster/Projects/APIAutomation/dataFiles/postMultipleStudents.json')
    json_request = json.loads(file.read())

    # Excel Code
    work_book = openpyxl.load_workbook('/home/qamaster/Projects/APIAutomation/dataFiles/studentsData.xlsx')
    sheet = work_book['Sheet']
    rows = sheet.max_row

    for i in range(2,rows+1):
        cell_first_name = sheet.cell(row=i,column=1)
        cell_middle_name = sheet.cell(row=i, column=2)
        cell_last_name = sheet.cell(row=i, column=3)
        cell_date_of_birth = sheet.cell(row=i, column=4)

        json_request['first_name'] = cell_first_name.value
        json_request['middle_name'] = cell_middle_name.value
        json_request['last_name'] = cell_last_name.value
        json_request['date_of_birth'] = cell_date_of_birth.value

        response = requests.post(API_URL, json_request)

        print(response.text)
        print(response.status_code)
        assert response.status_code == 201